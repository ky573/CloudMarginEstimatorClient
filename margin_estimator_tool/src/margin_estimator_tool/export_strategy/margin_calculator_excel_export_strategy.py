"""This module defines a strategy for exporting margin data to Excel format."""

import os
from typing import Dict, Any, List, Set, Tuple
from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet
from margin_estimator_tool.core.utils import flatten_dict
from margin_estimator_tool.export_strategy.export_strategy import (
    ExportStrategy,
)


class MarginCalculatorExcelExportStrategy(ExportStrategy):
    """Concrete strategy for exporting margin data to Excel."""

    def export(
        self, date: str, version: bool, data: List[Dict[str, Any]], output_path: str
    ) -> bool:
        """
        Concrete implementation for exporting margin data into Excel.

        Args:
            date: date to be included in file name
            version: version to be included in file name
            data: margin data to be exported
            output_path: directory where data will be exported

        Returns:
            True if there were any data to export, false otherwise
        """
        if not data:
            return False

        wb = Workbook()

        if "Sheet" in wb.sheetnames:
            del wb["Sheet"]

        ws1 = wb.create_sheet(title="portfolio_margin")
        ws2 = wb.create_sheet(title="drilldowns")

        portfolio_headers, drilldown_headers = self._prepare_headers(data)

        self._populate_sheet(ws1, portfolio_headers, data, "portfolio_margin")
        self._populate_sheet(ws2, drilldown_headers, data, "drilldowns")

        version_path = "LIVE" if version else "SOD"
        out_path = f"{date}_{version_path}_{self.type}.xlsx"
        file_path = os.path.join(output_path, out_path)
        wb.save(file_path)

        return True

    def _prepare_headers(
        self, margin_details: List[Dict[str, Any]]
    ) -> Tuple[List[str], List[str]]:
        """Prepares headers for portfolio and drilldowns sheets."""
        portfolio_headers = self._extract_headers(margin_details, "portfolio_margin")
        drilldown_headers = self._extract_headers(margin_details, "drilldowns")
        return portfolio_headers, drilldown_headers

    @staticmethod
    def _extract_headers(margin_details: List[Dict[str, Any]], key: str) -> List[str]:
        """Extracts headers by examining all data entries for the given key."""
        headers: Set[str] = set()
        for detail in margin_details:
            for item in detail[key]:
                flat_item = flatten_dict(item)
                headers.update(flat_item.keys())
        headers_list = ["business_date"] + sorted(list(headers))
        return headers_list

    @staticmethod
    def _populate_sheet(
        ws: Worksheet,
        headers: List[str],
        margin_details: List[Dict[str, Any]],
        data_key: str,
    ) -> None:
        """Populates a given worksheet with data based on the specified headers."""
        ws.append(headers)

        for detail in margin_details:
            business_date = detail.get("business_date", "")
            for item in detail[data_key]:
                flat_item = flatten_dict(item)
                row = [business_date] + [
                    flat_item.get(header, "") for header in headers[1:]
                ]
                ws.append(row)
